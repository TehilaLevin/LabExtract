"""Excel export: Hebrew headers, RTL sheet, text cells, orange flagged cells, column widths."""
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from src import excel_exporter as xe
from src.constants import FIELD_HEADERS, FLAG_COLOR, REVIEW_STATUS_HEADER, REVIEW_STATUS_KEY


def _df():
    return pd.DataFrame([
        {"first_name": "ישראל", "id_number": "000000018", "central_numbers": "937309, 937310",
         "received_date": "01/10/2026", REVIEW_STATUS_KEY: "אושר"},
        {"first_name": "", "id_number": None, "central_numbers": float("nan"),
         "received_date": "", REVIEW_STATUS_KEY: "לא אושר"},
    ])


def _sheet(data: bytes):
    return load_workbook(BytesIO(data)).active


def test_headers_are_hebrew_and_sheet_is_rtl():
    ws = _sheet(xe.to_xlsx_bytes(_df()))
    assert [c.value for c in ws[1]] == [FIELD_HEADERS["first_name"], FIELD_HEADERS["id_number"],
                                        FIELD_HEADERS["central_numbers"], FIELD_HEADERS["received_date"],
                                        REVIEW_STATUS_HEADER]
    assert ws.sheet_view.rightToLeft is True
    assert ws.max_row == 3  # header + one row per form


def test_values_are_text_and_keep_leading_zeros():
    ws = _sheet(xe.to_xlsx_bytes(_df()))
    assert ws["B2"].value == "000000018" and ws["B2"].data_type == "s"
    assert ws["D2"].value == "01/10/2026" and ws["D2"].data_type == "s"  # never turned into a date
    assert ws["C2"].value == "937309, 937310"
    assert ws["B3"].value is None and ws["C3"].value is None  # None / NaN -> empty cell


def test_only_highlighted_cells_are_orange():
    ws = _sheet(xe.to_xlsx_bytes(_df(), {(1, "received_date"), (1, REVIEW_STATUS_KEY)}))
    orange = {c.coordinate for row in ws.iter_rows(min_row=2) for c in row
              if c.fill.fill_type == "solid" and c.fill.start_color.rgb.endswith(FLAG_COLOR.lstrip("#"))}
    assert orange == {"D3", "E3"}


def test_column_width_follows_content_within_bounds():
    df = pd.DataFrame([{"first_name": "א" * 200, "id_number": "1"}])
    ws = _sheet(xe.to_xlsx_bytes(df))
    assert ws.column_dimensions["A"].width == xe.MAX_WIDTH
    assert xe.MIN_WIDTH <= ws.column_dimensions["B"].width < xe.MAX_WIDTH


def test_empty_dataframe_gives_header_only():
    ws = _sheet(xe.to_xlsx_bytes(pd.DataFrame(columns=["first_name", "id_number"])))
    assert ws.max_row == 1
