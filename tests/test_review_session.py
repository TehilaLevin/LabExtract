"""Review session: reviewer edits, approval rules, export rows and highlighted cells."""
from datetime import date

import pytest

from src import record_builder as rb
from src import review_session as rs
from src.constants import FIELD_KEYS, REVIEW_STATUS_APPROVED, REVIEW_STATUS_KEY, REVIEW_STATUS_PENDING

TODAY = date(2026, 10, 13)
VALID_ID = "000000018"


def _record(index=1, flags=None, suggestions=None, **values):
    base = {"first_name": "ישראל", "last_name": "ישראלי", "id_number": VALID_ID, "central_numbers": "937309",
            "received_date": "01/10/2026", "sending_entity": "מכבי", "commitment_numbers": "50364186"}
    base.update(values)
    return rb.Record(index, base, dict(flags or {}), {k: 0.9 for k in FIELD_KEYS},
                     {"source_file": "batch.pdf", "page_numbers": str(index), "form_type": "מכבי – מעבדה שולחת"},
                     suggestions=dict(suggestions or {}))


def _review(**kwargs):
    return rs.FormReview(_record(**kwargs))


# ---- values ------------------------------------------------------------------------------

def test_value_is_the_engine_value_until_the_reviewer_saves_one():
    review = _review()
    assert review.value("id_number") == VALID_ID
    review.edits["id_number"] = "000000026"
    assert review.value("id_number") == "000000026"
    assert review.record.values["id_number"] == VALID_ID  # the record is never changed


def test_suggestion_is_never_a_value():
    review = _review(first_name="", last_name="", flags={"first_name": "name_unverified"},
                     suggestions={"first_name": "ישראל"})
    assert review.value("first_name") == ""
    df, _ = rs.export_frame([review])
    assert df.loc[0, "first_name"] == ""


# ---- normalize_input ---------------------------------------------------------------------

@pytest.mark.parametrize("key, text, expected", [
    ("id_number", " 0000 00018 ", (VALID_ID, None)),
    ("id_number", "000000019", ("000000019", "invalid_id")),
    ("id_number", "", ("", "missing")),
    ("first_name", "  ישראל  ", ("ישראל", None)),
    ("first_name", "ישראל2", ("ישראל2", "invalid_format")),
    ("central_numbers", "937309; 937310 ,937309", ("937309, 937310", None)),
    ("central_numbers", "93730", ("93730", "invalid_format")),
    ("commitment_numbers", "X132767, 50364186", ("X132767, 50364186", None)),
    ("commitment_numbers", "5036-4186", ("5036-4186", "invalid_format")),
    ("received_date", "1.10.26", ("01/10/2026", None)),
    ("received_date", "2026-10-01", ("01/10/2026", None)),
    ("received_date", "31/02/2026", ("31/02/2026", "invalid_date")),
    ("received_date", "01/12/2026", ("01/12/2026", "date_out_of_range")),   # in the future
    ("sending_entity", "שירותי בריאות כללית", ("כללית", None)),
    ("sending_entity", "מרפאה פרטית כלשהי", ("מרפאה פרטית כלשהי", "sender_not_in_list")),
])
def test_normalize_input(key, text, expected):
    assert rs.normalize_input(key, text, today=TODAY) == expected


# ---- approval ----------------------------------------------------------------------------

def test_clean_form_is_approved_and_flags_close():
    review = _review(flags={"id_number": "low_confidence"})
    result = rs.approve(review, {}, today=TODAY)
    assert result.approved and result.warnings == {} and result.empty == []
    assert review.open_flags() == {}


def test_invalid_value_needs_approve_anyway():
    review = _review(id_number="000000019", flags={"id_number": "invalid_id"})
    result = rs.approve(review, {}, today=TODAY)
    assert not result.approved and result.warnings == {"id_number": "invalid_id"}
    assert review.open_flags() == {"id_number": "invalid_id"}
    assert rs.approve(review, {}, force=True, today=TODAY).approved


def test_reviewer_correction_fixes_the_warning():
    review = _review(id_number="000000019", flags={"id_number": "invalid_id"})
    result = rs.approve(review, {"id_number": "000000026"}, today=TODAY)
    assert result.approved and review.value("id_number") == "000000026"


def test_rejected_approval_still_keeps_what_the_reviewer_typed():
    review = _review()
    rs.approve(review, {"id_number": "123", "first_name": "יעל"}, today=TODAY)
    assert not review.approved
    assert review.value("id_number") == "123" and review.value("first_name") == "יעל"


def test_empty_fields_are_reported_but_never_block():
    review = _review(received_date="", flags={"received_date": "missing"})
    result = rs.approve(review, {}, today=TODAY)
    assert result.approved and result.empty == ["received_date"]


def test_approval_normalizes_values():
    review = _review()
    rs.approve(review, {"received_date": "1/10/26", "central_numbers": "937309;937310"}, today=TODAY)
    assert review.value("received_date") == "01/10/2026"
    assert review.value("central_numbers") == "937309, 937310"


def test_reopen():
    review = _review(flags={"id_number": "low_confidence"})
    rs.approve(review, {}, today=TODAY)
    rs.reopen(review)
    assert not review.approved and review.open_flags() == {"id_number": "low_confidence"}


# ---- navigation / summary -----------------------------------------------------------------

def test_next_unapproved_wraps_and_ends():
    reviews = rs.new_session([_record(i) for i in range(1, 5)])
    reviews[2].approved = True
    assert rs.next_unapproved(reviews, 1) == 3
    assert rs.next_unapproved(reviews, 3) == 0
    for r in reviews:
        r.approved = True
    assert rs.next_unapproved(reviews, 0) is None
    assert rs.unapproved(reviews) == []


def test_summary_frame():
    rec = _record(2, flags={"id_number": "invalid_id", "received_date": "missing"})
    rec.split_flags = ["split_uncertain"]
    df = rs.summary_frame(rs.new_session([_record(1), rec]))
    assert list(df["flagged_fields"]) == [0, 2]
    assert list(df["form_warnings"]) == [0, 1]
    assert list(df[REVIEW_STATUS_KEY]) == [REVIEW_STATUS_PENDING] * 2


# ---- export ------------------------------------------------------------------------------

def test_export_one_row_per_form_with_status_and_highlight():
    reviews = rs.new_session([_record(1, flags={"id_number": "low_confidence"}),
                              _record(2, flags={"received_date": "missing"}, received_date="")])
    rs.approve(reviews[0], {"id_number": "000000026"}, today=TODAY)
    df, highlight = rs.export_frame(reviews)
    assert len(df) == 2
    assert list(df.columns) == [*FIELD_KEYS, "source_file", "page_numbers", "form_type", REVIEW_STATUS_KEY]
    assert df.loc[0, "id_number"] == "000000026"
    assert list(df[REVIEW_STATUS_KEY]) == [REVIEW_STATUS_APPROVED, REVIEW_STATUS_PENDING]
    # approved form: nothing highlighted; unapproved form: its flagged field and its status
    assert highlight == {(1, "received_date"), (1, REVIEW_STATUS_KEY)}


def test_export_one_row_per_commitment():
    """3 central numbers + 2 commitment numbers -> 2 rows; the central numbers stay in one cell, the
    single-value fields repeat, the ID keeps its leading zeros."""
    review = rs.FormReview(_record(1, flags={"id_number": "low_confidence"}, id_number="000000018",
                                   central_numbers="937309, 937310, 937311",
                                   commitment_numbers="50364186, 50364187"))
    df, highlight = rs.export_frame([review])
    assert len(df) == 2
    assert list(df["commitment_numbers"]) == ["50364186", "50364187"]
    assert set(df["central_numbers"]) == {"937309, 937310, 937311"}
    assert set(df["id_number"]) == {"000000018"} and set(df["last_name"]) == {"ישראלי"}
    assert set(df[REVIEW_STATUS_KEY]) == {REVIEW_STATUS_PENDING}
    # the flag follows the form to both rows
    assert {(0, "id_number"), (1, "id_number"), (0, REVIEW_STATUS_KEY), (1, REVIEW_STATUS_KEY)} <= highlight


def test_export_form_without_commitment_gives_one_flagged_row():
    pending = rs.FormReview(_record(1, commitment_numbers=""))
    approved = rs.FormReview(_record(2, commitment_numbers=""))
    approved.approved = True
    df, highlight = rs.export_frame([pending, approved])
    assert len(df) == 2 and list(df["commitment_numbers"]) == ["", ""]
    assert (0, "commitment_numbers") in highlight          # empty and flagged while not approved
    assert (1, "commitment_numbers") not in highlight      # approved: the reviewer accepted the empty cell


def test_export_cells_after_split_are_text_with_leading_zeros():
    from io import BytesIO
    from openpyxl import load_workbook
    from src.excel_exporter import to_xlsx_bytes

    review = rs.FormReview(_record(1, id_number="001122334", commitment_numbers="007, 008"))
    df, highlight = rs.export_frame([review])
    ws = load_workbook(BytesIO(to_xlsx_bytes(df, highlight))).active
    assert ws.max_row == 3                                  # header + 2 commitments
    cells = {c.value for row in ws.iter_rows(min_row=2) for c in row}
    assert {"001122334", "007", "008"} <= cells
