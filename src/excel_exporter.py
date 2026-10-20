"""DataFrame -> in-memory .xlsx: Hebrew headers, right-to-left sheet, auto column width, orange flagged cells.

Nothing is written to disk: the workbook is built in a BytesIO and handed to the download button.
"""
from io import BytesIO
from typing import Iterable, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.constants import (
    ALL_HEADERS,
    FLAG_COLOR,
        REVIEW_STATUS_APPROVED,
    REVIEW_STATUS_HEADER,
    REVIEW_STATUS_KEY,
)
from src.validators import normalize_number_list

SHEET_TITLE = "התחייבויות"
HEADERS = {**ALL_HEADERS, REVIEW_STATUS_KEY: REVIEW_STATUS_HEADER}
MIN_WIDTH = 8
MAX_WIDTH = 60
_FILL = PatternFill(fill_type="solid", start_color=FLAG_COLOR.lstrip("#"), end_color=FLAG_COLOR.lstrip("#"))


def expand_per_commitment(frame: pd.DataFrame, highlight: Optional[Iterable] = None) -> tuple:
    """(frame, highlight) with ONE ROW PER COMMITMENT number, as the challenge requires.

    Single-value fields, metadata and the review status repeat on every row; all central numbers of the form
    stay together in one cell (joined with ", "). A form without a commitment number still gives one row, with
    an empty commitment cell that is flagged (orange) while the form is not approved. Highlights follow their
    form to every one of its rows.
    """
    highlight = set(highlight or ())
    rows, new_highlight = [], set()
    for pos, record in enumerate(frame.to_dict("records")):
        raw = record.get("commitment_numbers")
        numbers = normalize_number_list(raw if isinstance(raw, (str, list, tuple)) else None) or [""]
        flagged = {key for row, key in highlight if row == pos}
        if not numbers[0] and record.get(REVIEW_STATUS_KEY) != REVIEW_STATUS_APPROVED:
            flagged.add("commitment_numbers")
        for number in numbers:
            new_highlight |= {(len(rows), key) for key in flagged}
            rows.append({**record, "commitment_numbers": number})
    return pd.DataFrame(rows, columns=frame.columns), new_highlight


def to_xlsx_bytes(df: pd.DataFrame, highlight: Optional[Iterable] = None) -> bytes:
    """One row per DataFrame row, one column per DataFrame column (English keys -> Hebrew headers).

    highlight: (row position, column key) cells to fill orange. Every cell is written as text, so IDs and
    numbers keep their leading zeros and Excel never turns a value into a date or a number.
    """
    highlight = set(highlight or ())
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_TITLE
    ws.sheet_view.rightToLeft = True

    columns = list(df.columns)
    ws.append([HEADERS.get(c, c) for c in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for pos, row in enumerate(df.itertuples(index=False)):
        for col, (key, value) in enumerate(zip(columns, row), start=1):
            cell = ws.cell(row=pos + 2, column=col, value=_text(value))
            cell.number_format = "@"
            cell.alignment = Alignment(horizontal="right")
            if (pos, key) in highlight:
                cell.fill = _FILL
    ws.freeze_panes = "A2"

    for col, key in enumerate(columns, start=1):
        texts = [HEADERS.get(key, key)] + [_text(v) for v in df[key]]
        width = max(len(t) for t in texts) + 2
        ws.column_dimensions[get_column_letter(col)].width = max(MIN_WIDTH, min(MAX_WIDTH, width))

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _text(value) -> str:
    if value is None or (isinstance(value, float) and value != value):  # None / NaN
        return ""
    return str(value)
